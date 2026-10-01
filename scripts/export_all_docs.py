"""
=============================================================================
BỘ KỊCH BẢN XUẤT TOÀN BỘ BÁO CÁO SANG MICROSOFT WORD & PDF
Nhóm 11 - CSE703093 / CSE703153 (ThS. Vũ Quang Dũng - ĐH Phenikaa)
=============================================================================
Chuyển đổi 2 tài liệu Markdown sang DOCX và PDF thông qua Microsoft Word COM:
  1. BaoCao_BTL_Nhom11 (.docx, .pdf)
  2. HUONG_DAN_KIEM_TRA_DU_AN (.docx, .pdf)
=============================================================================
"""

import os
import sys
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(__file__).resolve().parent.parent

DOCUMENTS = [
    {
        "md": PROJECT_ROOT / "report" / "BaoCao_BTL_Nhom11.md",
        "docx": PROJECT_ROOT / "report" / "BaoCao_BTL_Nhom11.docx",
        "pdf": PROJECT_ROOT / "report" / "BaoCao_BTL_Nhom11.pdf",
        "title": "Báo cáo BTL Chính thức"
    },
    {
        "md": PROJECT_ROOT / "HUONG_DAN_KIEM_TRA_DU_AN.md",
        "docx": PROJECT_ROOT / "HUONG_DAN_KIEM_TRA_DU_AN.docx",
        "pdf": PROJECT_ROOT / "HUONG_DAN_KIEM_TRA_DU_AN.pdf",
        "title": "Hướng dẫn Kiểm tra Dự án"
    }
]

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def style_table(table, header_bg="1F4E79"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            if i == 0:
                set_cell_background(cell, header_bg)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(9.5)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F2F5F8")
                else:
                    set_cell_background(cell, "FFFFFF")
                for run in cell.paragraphs[0].runs:
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(9)

def convert_md_to_docx(md_path, docx_path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_data = []

    for line in lines:
        raw_line = line.strip()

        if raw_line.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                code_text = "\n".join(code_lines)
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 30, 80)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line.rstrip())
            continue

        # Image handling: ![caption](path)
        if raw_line.startswith("![") and "](" in raw_line and raw_line.endswith(")"):
            caption = raw_line[2:raw_line.find("](")]
            rel_img_path = raw_line[raw_line.find("](") + 2 : -1]
            full_img_path = (PROJECT_ROOT / "results" / rel_img_path).resolve()
            if not full_img_path.exists():
                full_img_path = (PROJECT_ROOT / rel_img_path).resolve()
            if full_img_path.exists():
                try:
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(2)
                    run_img = p_img.add_run()
                    run_img.add_picture(str(full_img_path), width=Inches(5.8))
                    
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(8)
                    r_cap = p_cap.add_run(f"Hình: {caption}")
                    r_cap.font.name = "Segoe UI"
                    r_cap.font.size = Pt(9)
                    r_cap.font.italic = True
                    r_cap.font.color.rgb = RGBColor(80, 80, 80)
                except Exception as e:
                    print(f"[-] Lỗi chèn ảnh {full_img_path}: {e}")
            continue

        if raw_line.startswith("|") and raw_line.endswith("|"):
            cols = [c.strip() for c in raw_line[1:-1].split("|")]
            if all(c.startswith(":") or c.startswith("-") for c in cols if c):
                continue
            table_data.append(cols)
            in_table = True
            continue
        else:
            if in_table and table_data:
                num_rows = len(table_data)
                num_cols = max(len(r) for r in table_data)
                t = doc.add_table(rows=num_rows, cols=num_cols)
                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        if c_idx < num_cols:
                            clean_val = val.replace("**", "").replace("*", "").replace("<br>", "\n")
                            t.cell(r_idx, c_idx).text = clean_val
                style_table(t)
                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(6)
                table_data = []
                in_table = False

        if not raw_line:
            continue

        if raw_line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(raw_line[2:])
            r.font.name = "Segoe UI"
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor(31, 78, 121)
        elif raw_line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(raw_line[3:])
            r.font.name = "Segoe UI"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(46, 117, 182)
        elif raw_line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(raw_line[4:])
            r.font.name = "Segoe UI"
            r.font.size = Pt(11.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(68, 84, 106)
        elif raw_line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(raw_line[5:])
            r.font.name = "Segoe UI"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(80, 80, 80)
        elif raw_line.startswith("- ") or raw_line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            clean_text = raw_line[2:].replace("**", "").replace("*", "")
            r = p.add_run(clean_text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
        elif raw_line.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(raw_line[2:])
            r.font.name = "Segoe UI"
            r.font.size = Pt(10)
            r.font.italic = True
            r.font.color.rgb = RGBColor(50, 50, 100)
            p.paragraph_format.left_indent = Inches(0.4)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)
            clean_text = raw_line.replace("**", "").replace("*", "")
            r = p.add_run(clean_text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(10.5)

    doc.save(str(docx_path))
    print(f"  [+] Đã xuất Word: {docx_path.name}")

def convert_docx_to_pdf(docx_path, pdf_path, word_app):
    wdoc = word_app.Documents.Open(str(docx_path))
    wdoc.SaveAs(str(pdf_path), FileFormat=17) # 17 = wdFormatPDF
    wdoc.Close()
    print(f"  [+] Đã xuất PDF : {pdf_path.name}")

def main():
    print("\n" + "="*70)
    print("  TIẾN TRÌNH XUẤT TÀI LIỆU BÁO CÁO BTL NHÓM 11 SANG WORD VÀ PDF")
    print("="*70)
    
    # Bước 1: Sinh file Word
    print("\n[*] Bước 1: Biên dịch Markdown sang Microsoft Word (.docx)...")
    for d in DOCUMENTS:
        if d["md"].exists():
            convert_md_to_docx(d["md"], d["docx"])
            
    # Bước 2: Chuyển đổi sang PDF qua Microsoft Word COM
    print("\n[*] Bước 2: Khởi tạo Microsoft Word COM Automation để xuất PDF...")
    try:
        import win32com.client
        word_app = win32com.client.Dispatch('Word.Application')
        word_app.Visible = False
        try:
            for d in DOCUMENTS:
                if d["docx"].exists():
                    convert_docx_to_pdf(d["docx"], d["pdf"], word_app)
        finally:
            word_app.Quit()
        print("\n[+] Toàn bộ tài liệu PDF đã được biên dịch hoàn tất 100%!")
    except Exception as e:
        print(f"[-] Lỗi khi xuất PDF qua Word COM: {e}")
        
    print("="*70 + "\n")

if __name__ == "__main__":
    main()
